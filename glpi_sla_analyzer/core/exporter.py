"""Выгрузка отчёта в xlsx, csv и json."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_HEADERS = [
    "ID",
    "Заголовок",
    "Статус",
    "Категория",
    "Начало",
    "Конец",
    "Pending, ч",
    "Источник Pending",
    "Чистое время, ч",
    "solve_delay_stat, ч",
    "Превышение SLA, ч",
    "Предупреждение",
]


_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _fmt(value):
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if value is None:
        return ""
    if isinstance(value, str):
        return _ILLEGAL.sub("", value)
    return value


def ticket_rows(rows: list[dict]) -> list[list]:
    out = []
    for row in rows:
        out.append(
            [
                row.get("id"),
                _fmt(row.get("title")),
                row.get("status"),
                _fmt(row.get("category")),
                _fmt(row.get("opened")),
                _fmt(row.get("end")),
                row.get("pending_hours"),
                row.get("pending_source"),
                row.get("net_hours"),
                row.get("solve_delay_stat_hours"),
                row.get("excess_hours") if row.get("included") else "",
                _fmt(row.get("warning") or ("Pending не вычтен" if row.get("pending_not_deducted") else "")),
            ]
        )
    return out


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(_HEADERS)
        writer.writerows(ticket_rows(rows))


def write_json(path: Path, summary: dict, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"summary": summary, "tickets": ticket_rows(rows)}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=_fmt), encoding="utf-8")


def write_xlsx(path: Path, summary: dict, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Сводка"
    lines = [
        ("Заявок в среднем", summary.get("n")),
        ("Незавершённые", summary.get("open_count")),
        ("Пропущены (ошибка дат)", summary.get("skipped")),
        ("Среднее, ч", summary.get("avg_hours")),
        ("Превысили SLA", summary.get("over_count")),
        ("Доля превышений, %", summary.get("over_share_pct")),
        ("Среднее превышение, ч", summary.get("avg_excess")),
        ("Макс. превышение, ч", summary.get("max_excess")),
        ("ID рекордсмена", summary.get("max_excess_id")),
        ("Медиана", summary.get("median")),
        ("P90", summary.get("p90")),
        ("P95", summary.get("p95")),
        ("Минимум", summary.get("min_hours")),
        ("Максимум", summary.get("max_hours")),
    ]
    for idx, (label, value) in enumerate(lines, start=1):
        sheet.cell(idx, 1, label)
        sheet.cell(idx, 2, value if value is not None else "")
    sheet.cell(len(lines) + 2, 1, "Неделя")
    sheet.cell(len(lines) + 2, 2, "Штук")
    sheet.cell(len(lines) + 2, 3, "Среднее, ч")
    sheet.cell(len(lines) + 2, 4, "Превысили SLA")
    row_i = len(lines) + 3
    for item in summary.get("weeks") or []:
        sheet.cell(row_i, 1, item["period"])
        sheet.cell(row_i, 2, item["count"])
        sheet.cell(row_i, 3, item["avg_hours"])
        sheet.cell(row_i, 4, item["over_sla"])
        row_i += 1

    table = book.create_sheet("Заявки")
    red = PatternFill("solid", fgColor="F4C7C3")
    for col, title in enumerate(_HEADERS, start=1):
        cell = table.cell(1, col, title)
        cell.font = Font(bold=True)
    for r, values in enumerate(ticket_rows(rows), start=2):
        for col, value in enumerate(values, start=1):
            table.cell(r, col, value)
        source = rows[r - 2]
        if source.get("over_sla"):
            for col in range(1, len(_HEADERS) + 1):
                table.cell(r, col).fill = red
    book.save(path)
